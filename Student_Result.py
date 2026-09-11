# Student Result Management System

from openpyxl import Workbook, load_workbook
import os


# Calculate Result
def calculate_result(marks):

    total = sum(marks)
    percentage = total / 5

    if percentage >= 75:
        grade = "A"
    elif percentage >= 60:
        grade = "B"
    elif percentage >= 50:
        grade = "C"
    elif percentage >= 35:
        grade = "D"
    else:
        grade = "F"

    if percentage >= 35:
        status = "PASS"
    else:
        status = "FAIL"

    return total, percentage, grade, status


# Add Student
def add_student():

    print("\n----- Add Student Result -----")

    roll_no = int(input("Enter Roll No: "))
    name = input("Enter Student Name: ")
    student_class = input("Enter Class: ")

    print("\nEnter Marks:")

    s1 = int(input("Marathi: "))
    s2 = int(input("Hindi: "))
    s3 = int(input("English: "))
    s4 = int(input("Science: "))
    s5 = int(input("History: "))

    marks = [s1, s2, s3, s4, s5]

    total, percentage, grade, status = calculate_result(marks)

    print("\n----- Student Result -----")
    print("Roll No    :", roll_no)
    print("Name       :", name)
    print("Class      :", student_class)
    print("Total      :", total)
    print("Percentage :", percentage)
    print("Grade      :", grade)
    print("Status     :", status)

    # Excel file
    file_name = "student_results.xlsx"

    if os.path.exists(file_name):

        workbook = load_workbook(file_name)
        sheet = workbook.active

    else:

        workbook = Workbook()
        sheet = workbook.active

        # Excel headings
        sheet.append([
            "Roll No",
            "Name",
            "Class",
            "Marathi",
            "Hindi",
            "English",
            "Science",
            "History",
            "Total",
            "Percentage",
            "Grade",
            "Status"
        ])

    # Add student data
    sheet.append([
        roll_no,
        name,
        student_class,
        s1,
        s2,
        s3,
        s4,
        s5,
        total,
        percentage,
        grade,
        status
    ])

    workbook.save(file_name)

    print("\nStudent result saved successfully!")


# Get Student Result
def get_result():

    print("\n----- Get Student Result -----")

    roll_no = int(input("Enter Roll No: "))

    file_name = "student_results.xlsx"

    if not os.path.exists(file_name):

        print("No student data found.")
        return

    workbook = load_workbook(file_name)
    sheet = workbook.active

    found = False

    for row in sheet.iter_rows(min_row=2, values_only=True):

        if row[0] == roll_no:

            print("\n------------------------------")
            print("Student Result")
            print("------------------------------")

            print("Roll No     :", row[0])
            print("Name        :", row[1])
            print("Class       :", row[2])
            print("Total       :", row[8])
            print("Percentage  :", row[9])
            print("Grade       :", row[10])
            print("Status      :", row[11])

            print("------------------------------")

            found = True
            break

    if found == False:

        print("Student not found.")


# Show All Student Data
def show_all_data():

    print("\n----- All Student Data -----")

    file_name = "student_results.xlsx"

    if not os.path.exists(file_name):

        print("No student data found.")
        return

    workbook = load_workbook(file_name)
    sheet = workbook.active

    print("\nRoll No | Name | Class | Total | Percentage | Grade | Status")
    print("--------------------------------------------------------------")

    for row in sheet.iter_rows(min_row=2, values_only=True):

        print(
            row[0], "|",
            row[1], "|",
            row[2], "|",
            row[8], "|",
            row[9], "|",
            row[10], "|",
            row[11]
        )


# Menu
def menu():

    while True:

        print("\n================================")
        print("    STUDENT RESULT MANAGEMENT")
        print("================================")

        print("1. Add Student Result")
        print("2. Get Student Result")
        print("3. Show All Student Data")
        print("4. Exit")

        choice = int(input("Enter your choice: "))

        if choice == 1:

            add_student()

        elif choice == 2:

            get_result()

        elif choice == 3:

            show_all_data()

        elif choice == 4:

            print("Thank you!")
            break

        else:

            print("Please enter a valid choice.")


# Start Program
menu()
